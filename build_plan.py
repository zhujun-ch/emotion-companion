from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.2)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = '宋体'
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    s.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.space_after = Pt(5)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name, size in [('Title', 18), ('Heading 1', 14), ('Heading 2', 12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(6)
table = doc.add_table(rows=1, cols=1)
table.autofit = False
table.columns[0].width = Cm(16.6)
props = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
    el = OxmlElement('w:' + edge)
    el.set(qn('w:val'), 'single' if edge in ['top', 'left', 'bottom', 'right'] else 'nil')
    el.set(qn('w:sz'), '6')
    el.set(qn('w:color'), '000000')
    borders.append(el)
props.append(borders)
cell = table.cell(0, 0)
margins = OxmlElement('w:tcMar')
for edge in ['top', 'left', 'bottom', 'right']:
    el = OxmlElement('w:' + edge)
    el.set(qn('w:w'), '100')
    el.set(qn('w:type'), 'dxa')
    margins.append(el)
cell._tc.get_or_add_tcPr().append(margins)
first = True
for line in (ROOT / 'PROJECT_PLAN.md').read_text(encoding='utf-8').splitlines():
    if not line.strip():
        continue
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    first = False
    if line.startswith('# '):
        p.style = 'Title'
        p.add_run(line[2:])
    elif line.startswith('## '):
        p.style = 'Heading 1'
        p.add_run(line[3:])
    else:
        p.add_run(line)
    p.paragraph_format.widow_control = True
# A single physical cell spans pages, with no fixed height or cantSplit flag.
footer = sec.footer.paragraphs[0]
footer.alignment = 1
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
output = ROOT / '中文情绪识别智能陪伴小程序_项目方案.docx'
doc.save(output)
print(output)
print('body_tables=1 rows=1 cells=1 nested_tables=0')
